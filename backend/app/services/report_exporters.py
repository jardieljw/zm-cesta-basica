import csv
import io
import datetime
import textwrap
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.schemas.reports import ReportResponse
from app.schemas.form_print import FormPrintOptions, column_weight, select_form_columns


def _format_currency_ptbr(cents: int) -> str:
    """Formats cents as Brazilian Real string: R$ 1.234,56."""
    val = (cents or 0) / 100.0
    formatted = f"{val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {formatted}"


def export_report_to_csv(report: ReportResponse) -> bytes:
    """Exports a report to CSV using UTF-8-SIG (Excel compatible in pt-BR)."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";", quoting=csv.QUOTE_MINIMAL)

    # Header metadata
    writer.writerow([report.metadata.company_name])
    writer.writerow([report.metadata.title])
    writer.writerow([f"Período/Filtro: {report.metadata.filter_info.label}"])
    writer.writerow([f"Gerado em: {report.metadata.generated_at.strftime('%d/%m/%Y %H:%M:%S')}"])
    writer.writerow([])

    # Table headers
    writer.writerow(report.headers)

    # Data rows
    for row in report.rows:
        row_values = []
        for col in report.columns:
            val = row.get(col, "")
            if isinstance(val, (int, float)):
                if col.endswith("_cents") or "valor" in col or "preco" in col:
                    # Currency field
                    row_values.append(f"{(val / 100.0):.2f}".replace(".", ","))
                else:
                    row_values.append(str(val))
            elif isinstance(val, datetime.datetime):
                row_values.append(val.strftime("%d/%m/%Y %H:%M"))
            elif isinstance(val, datetime.date):
                row_values.append(val.strftime("%d/%m/%Y"))
            elif val is None:
                row_values.append("")
            else:
                row_values.append(str(val))
        writer.writerow(row_values)

    # Totals row if available
    if report.totals:
        writer.writerow([])
        totals_row = ["TOTAIS"]
        for col in report.columns[1:]:
            if col in report.totals:
                val = report.totals[col]
                if col.endswith("_cents"):
                    totals_row.append(f"{(val / 100.0):.2f}".replace(".", ","))
                else:
                    totals_row.append(str(val))
            else:
                totals_row.append("")
        writer.writerow(totals_row)

    return output.getvalue().encode("utf-8-sig")


def export_report_to_xlsx(report: ReportResponse) -> bytes:
    """Exports a report to a professionally formatted Excel spreadsheet using openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report.metadata.report_type[:30].replace("/", "-")

    # Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    meta_font = Font(name="Calibri", size=10, italic=True, color="595959")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F5F8", end_color="F2F5F8", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="000000")
    total_fill = PatternFill(start_color="E6EEF4", end_color="E6EEF4", fill_type="solid")
    
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )
    total_border = Border(
        top=Side(style="thin", color="1F497D"),
        bottom=Side(style="double", color="1F497D")
    )

    # Title block
    ws.append([report.metadata.company_name])
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=11, bold=True, color="595959")
    
    ws.append([report.metadata.title])
    ws.cell(row=2, column=1).font = title_font

    ws.append([f"Período: {report.metadata.filter_info.label} | Emissão: {report.metadata.generated_at.strftime('%d/%m/%Y %H:%M:%S')}"])
    ws.cell(row=3, column=1).font = meta_font
    ws.append([]) # Empty row

    # Headers
    header_row_idx = 5
    ws.append(report.headers)
    for col_idx in range(1, len(report.headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row_idx].height = 24

    # Data rows
    curr_row_idx = header_row_idx + 1
    for r_idx, row in enumerate(report.rows):
        row_vals = []
        for col in report.columns:
            val = row.get(col, "")
            row_vals.append(val)
        ws.append(row_vals)

        # Style and format cells
        is_even = (r_idx % 2 == 1)
        for col_idx, col in enumerate(report.columns, start=1):
            cell = ws.cell(row=curr_row_idx, column=col_idx)
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill

            val = cell.value
            if col.endswith("_cents") or "valor" in col or "preco" in col:
                if isinstance(val, (int, float)):
                    cell.value = val / 100.0
                    cell.number_format = '"R$" #,##0.00'
                    cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, int):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, datetime.datetime):
                cell.value = val.strftime("%d/%m/%Y %H:%M")
                cell.alignment = Alignment(horizontal="center")
            elif isinstance(val, datetime.date):
                cell.value = val.strftime("%d/%m/%Y")
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.alignment = Alignment(horizontal="left")

        curr_row_idx += 1

    # Totals Row
    if report.totals:
        totals_row_vals = ["TOTAIS"]
        for col in report.columns[1:]:
            if col in report.totals:
                totals_row_vals.append(report.totals[col])
            else:
                totals_row_vals.append("")
        ws.append(totals_row_vals)

        for col_idx, col in enumerate(report.columns, start=1):
            cell = ws.cell(row=curr_row_idx, column=col_idx)
            cell.font = total_font
            cell.fill = total_fill
            cell.border = total_border
            val = cell.value
            if col.endswith("_cents") or "valor" in col:
                if isinstance(val, (int, float)):
                    cell.value = val / 100.0
                    cell.number_format = '"R$" #,##0.00'
                    cell.alignment = Alignment(horizontal="right")
            elif isinstance(val, int):
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right")
        curr_row_idx += 1

    # Adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title row from max_len calculation
            if cell.row < 5:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_blank_template_to_xlsx(
    template_type: str, report: Optional[ReportResponse] = None,
    options: Optional[FormPrintOptions] = None,
) -> bytes:
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.page import PageMargins
    from app.services.reports import get_blank_form_data

    report = report or get_blank_form_data(template_type)
    blank = report.metadata.report_type.startswith('blank-')
    # Preserve the legacy GET export's row height/count; configurable exports share PDF defaults.
    options = options or FormPrintOptions(blank_row_count=20, row_height_mm=(28 if blank else 26) / (72 / 25.4))
    report = select_form_columns(report, options)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = {'payments': 'Pagamentos', 'payables': 'Contas a pagar', 'physical_inventory': 'Contagem física'}.get(template_type, 'Formulário')
    count = len(report.columns)
    last = get_column_letter(count)
    landscape = options.landscape(count)
    weights = [options.column_weights.get(c, column_weight(c)) for c in report.columns]
    printable_mm = (297 if landscape else 210) - options.margin_left_mm - options.margin_right_mm
    widths_mm = [printable_mm * weight / sum(weights) for weight in weights]
    for idx, text in enumerate([
        report.metadata.company_name, report.metadata.title,
        f"Consulta: {report.metadata.generated_at.astimezone(datetime.timezone(datetime.timedelta(hours=-3))).strftime('%d/%m/%Y %H:%M')} | {report.metadata.filter_info.label}",
    ], start=1):
        ws.cell(idx, 1, text)
        if count > 1:
            ws.merge_cells(start_row=idx, start_column=1, end_row=idx, end_column=count)
        ws.cell(idx, 1).font = Font(name='Calibri', size=11 if idx == 2 else 9, bold=idx < 3)
        ws.cell(idx, 1).alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[idx].height = {1: 18, 2: 32, 3: 22}[idx]
    ws.append([])
    ws.row_dimensions[4].height = 28 * 72 / 25.4 - 72
    for i, header in enumerate(report.headers, start=1):
        cell = ws.cell(5, i, header)
        cell.font = Font(name='Calibri', size=options.font_size, bold=True)
        cell.fill = PatternFill('solid', fgColor='E6E6E6')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[5].height = 18 * 72 / 25.4
    border = Border(**{side: Side(style='thin', color='555555') for side in ['left', 'right', 'top', 'bottom']})
    data = [{c: '' for c in report.columns} for _ in range(options.blank_rows(count))] if blank else report.rows
    if not blank and not data:
        ws.cell(6, 1, 'Nenhum registro encontrado.')
    for row_idx, row in enumerate(data, start=6):
        ws.row_dimensions[row_idx].height = options.row_height_mm * 72 / 25.4
        for i, col in enumerate(report.columns, start=1):
            value = row.get(col, '')
            cell = ws.cell(row_idx, i)
            # Force strings to text: descriptions must never become spreadsheet formulas.
            if isinstance(value, str):
                try:
                    if col.endswith('_date') and value:
                        value = datetime.datetime.fromisoformat(value.replace('Z', '+00:00'))
                except ValueError:
                    pass
            if isinstance(value, datetime.datetime):
                if col == 'due_date':
                    value = value.replace(tzinfo=None)
                elif value.tzinfo:
                    value = value.astimezone(datetime.timezone(datetime.timedelta(hours=-3))).replace(tzinfo=None)
                else:
                    value = value.replace(tzinfo=datetime.timezone.utc).astimezone(datetime.timezone(datetime.timedelta(hours=-3))).replace(tzinfo=None)
                cell.value = value
                cell.number_format = 'dd/mm/yyyy'
            elif isinstance(value, datetime.date):
                cell.value = value
                cell.number_format = 'dd/mm/yyyy'
            elif col.endswith('_cents') and isinstance(value, (int, float)):
                cell.value = value / 100
                cell.number_format = '"R$" #,##0.00'
            else:
                cell.value = value
                if isinstance(value, str):
                    cell.data_type = 's'
            cell.border = border
            cell.font = Font(name='Calibri', size=options.font_size)
            cell.alignment = Alignment(vertical='center', horizontal='right' if col.endswith('_cents') else 'left', wrap_text=True)
            if not blank:
                chars_per_line = max(1, int((widths_mm[i - 1] * 72 / 25.4 - 6) / (options.font_size * 0.6)))
                display = _format_currency_ptbr(row.get(col)) if col.endswith('_cents') and isinstance(row.get(col), (int, float)) else str(value or '')
                lines = sum(max(1, len(textwrap.wrap(part, width=chars_per_line))) for part in display.split('\n'))
                ws.row_dimensions[row_idx].height = max(ws.row_dimensions[row_idx].height, lines * options.font_size * 1.25 + 6)
    end_row = 5 + max(len(data), 1)
    if not blank and report.totals:
        end_row += 1
        for i, col in enumerate(report.columns, start=1):
            cell = ws.cell(end_row, i)
            if col in report.totals:
                cell.value = report.totals[col] / 100 if col.endswith('_cents') else report.totals[col]
                if col.endswith('_cents'):
                    cell.number_format = '"R$" #,##0.00'
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='E6E6E6')
    # Convert physical table width to Excel character units without shrinking the printed page.
    for i, weight in enumerate(weights, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(2, (printable_mm * 96 / 25.4 * weight / sum(weights) - 5) / 7)
    ws.freeze_panes = 'A6'
    ws.print_title_rows = '1:5'
    ws.print_options.horizontalCentered = True
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.scale = 100
    ws.page_margins = PageMargins(left=options.margin_left_mm / 25.4, right=options.margin_right_mm / 25.4,
        top=options.margin_top_mm / 25.4, bottom=options.margin_bottom_mm / 25.4, header=0.15, footer=0.15)
    ws.oddFooter.left.text = 'Responsável: __________________' if options.show_signatures else None
    ws.oddFooter.right.text = 'Conferente: __________________' if options.show_signatures else None
    ws.oddFooter.center.text = 'Página &P de &N'
    available_points = ((210 if landscape else 297) - options.margin_top_mm - options.margin_bottom_mm - 28 - 18 - (18 if options.show_signatures else 9)) * 72 / 25.4
    occupied = 0
    for row_idx in range(6, end_row + 1):
        height = ws.row_dimensions[row_idx].height or 15
        if occupied and occupied + height > available_points:
            ws.row_breaks.append(Break(id=row_idx - 1))
            occupied = 0
        occupied += height
    ws.print_area = f'A1:{last}{end_row}'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
