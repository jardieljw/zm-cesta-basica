import datetime
import io
import openpyxl
import pytest
from app.schemas.form_print import FormPrintOptions
from app.services.reports import get_blank_form_data
from app.models.financial import AccountPayable
from app.models.customers import Customer
from app.models.sales import Sale, SalesReceivable


@pytest.mark.parametrize('form_type', ['payments', 'payables', 'physical_inventory'])
def test_configured_export_selected_columns_and_print_dimensions(client, authenticated_device, form_type):
    headers = authenticated_device[2]
    snapshot = client.get(f'/api/v1/reports/blank-templates/{form_type}', headers=headers).json()
    columns = [snapshot['columns'][-1], snapshot['columns'][0]]
    result = client.post(f'/api/v1/reports/blank-templates/{form_type}/export', headers=headers, json={
        'report': snapshot, 'options': {'columns': columns, 'orientation': 'portrait', 'blank_row_count': 7,
            'margin_left_mm': 20, 'margin_top_mm': 18, 'row_height_mm': 12,
            'column_weights': {columns[0]: 3, columns[1]: 1}},
    })
    assert result.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(result.content)).active
    assert [ws.cell(5, i).value for i in [1, 2]] == [snapshot['headers'][snapshot['columns'].index(c)] for c in columns]
    assert ws.max_column == 2
    assert ws.page_setup.orientation == 'portrait'
    assert str(ws.page_setup.paperSize) == ws.PAPERSIZE_A4
    assert ws.print_title_rows == '$1:$5'
    assert ws.page_margins.left == pytest.approx(20 / 25.4)
    assert ws.page_margins.top == pytest.approx(18 / 25.4)
    assert [ws.row_dimensions[r].height for r in range(6, 13)] == pytest.approx([12 * 72 / 25.4] * 7)
    assert ws.row_dimensions[13].height is None
    assert ws.column_dimensions['A'].width > ws.column_dimensions['B'].width * 2


@pytest.mark.parametrize('options', [{'columns': []}, {'columns': ['qty', 'qty']}, {'columns': ['unknown']},
    {'row_height_mm': 0}, {'font_size': 2}, {'margin_left_mm': 99}, {'orientation': 'invalid'}, {'column_weights': {'qty': 0}}])
def test_invalid_export_configuration_is_rejected(client, authenticated_device, options):
    headers = authenticated_device[2]
    report = client.get('/api/v1/reports/blank-templates/payments', headers=headers).json()
    result = client.post('/api/v1/reports/blank-templates/payments/export', headers=headers, json={'report': report, 'options': options})
    assert result.status_code == 422


def test_cancelled_titles_excluded_from_rows_and_totals(test_session):
    now = datetime.datetime(2026, 10, 14)
    test_session.add(Customer(id='forms-customer', name='Cliente teste'))
    for index, status in enumerate(['OPEN', 'CANCELED']):
        sale = Sale(id=f'forms-sale-{index}', sale_number=9000+index, customer_id='forms-customer',
            total_amount_cents=10000, status='CONFIRMED' if status == 'OPEN' else 'CANCELED')
        test_session.add(sale)
        test_session.flush()
        test_session.add(SalesReceivable(id=f'forms-rec-{index}', sale_id=sale.id, customer_id='forms-customer',
            due_date=now, original_amount_cents=10000, paid_amount_cents=0, balance_cents=10000, status=status))
        test_session.add(AccountPayable(id=f'forms-pay-{index}', payable_number=9000+index, description='Conta',
            beneficiary='Fornecedor', category='OUTROS', due_date=now, original_amount_cents=20000,
            paid_amount_cents=3000, balance_cents=17000, status=status, notes='NF 123'))
    test_session.commit()
    rec = get_blank_form_data('payments', db=test_session, filled=True)
    pay = get_blank_form_data('payables', db=test_session, filled=True)
    assert len(rec.rows) == len(pay.rows) == 1
    assert rec.totals['total_cents'] == 10000
    assert pay.totals == {'amount_cents': 20000, 'paid_amount_cents': 3000, 'balance_cents': 17000}
    assert pay.rows[0]['notes'] == 'NF 123'
    assert pay.rows[0]['paid_amount_cents'] == 3000


def test_snapshot_export_does_not_query_new_business_data(client, authenticated_device):
    headers = authenticated_device[2]
    report = client.get('/api/v1/reports/blank-templates/payments?filled=true', headers=headers).json()
    report['rows'] = [{'customer': '=SUM(A1:A2)', 'total_cents': 12345}]
    report['metadata']['total_rows'] = 1
    report['totals'] = {'total_cents': 12345}
    result = client.post('/api/v1/reports/blank-templates/payments/export', headers=headers,
        json={'report': report, 'options': {'columns': ['customer', 'total_cents'], 'show_signatures': False}})
    assert result.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(result.content)).active
    assert ws['A6'].value == '=SUM(A1:A2)'
    assert ws['A6'].data_type == 's'
    assert ws['B6'].value == 123.45
    assert ws['B7'].value == 123.45
    assert ws.page_setup.orientation == 'portrait'
    assert 'Responsável' not in (ws.oddFooter.left.text or '')


def test_auto_rows_default_and_custom():
    assert FormPrintOptions().blank_rows(11) == 12
    assert FormPrintOptions().blank_rows(3) == 22
    assert FormPrintOptions(blank_row_count=37).blank_rows(3) == 37


def test_long_cells_expand_and_due_dates_preserve_calendar_day(client, authenticated_device):
    headers = authenticated_device[2]
    report = client.get('/api/v1/reports/blank-templates/payables?filled=true', headers=headers).json()
    report['rows'] = [{'description': 'Observação detalhada de conferência. ' * 15, 'due_date': '2026-10-14T00:00:00Z'}]
    report['metadata']['total_rows'] = 1
    result = client.post('/api/v1/reports/blank-templates/payables/export', headers=headers,
        json={'report': report, 'options': {'columns': ['description', 'due_date']}})
    assert result.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(result.content)).active
    assert ws.row_dimensions[6].height > 9 * 72 / 25.4
    assert ws['B6'].value.date() == datetime.date(2026, 10, 14)


def test_wrong_model_snapshot_is_rejected(client, authenticated_device):
    headers = authenticated_device[2]
    report = client.get('/api/v1/reports/blank-templates/payments', headers=headers).json()
    result = client.post('/api/v1/reports/blank-templates/payables/export', headers=headers,
        json={'report': report, 'options': {}})
    assert result.status_code == 422
